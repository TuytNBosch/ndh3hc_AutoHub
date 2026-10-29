"""Reads a Kick Off Meeting .pptx (optionally password protected) into project facts."""
import io
import logging
import os
from typing import Iterator

import msoffcrypto
from pptx import Presentation

from autohub.core.records import KickOffInfo
from autohub.services.KickOff.parser import parse_kick_off

_logger = logging.getLogger(__name__)

#: Tried in order when the deck is encrypted.
DEFAULT_PASSWORDS = ("honda", "hb")


def _decrypt(path: str, passwords: tuple[str, ...]) -> io.BytesIO | None:
    """The decrypted deck, or None when no password opens it (an unencrypted deck is returned as is)."""
    with open(path, "rb") as handle:
        data = handle.read()
    if not msoffcrypto.OfficeFile(io.BytesIO(data)).is_encrypted():
        return io.BytesIO(data)
    for password in passwords:
        office_file = msoffcrypto.OfficeFile(io.BytesIO(data))
        office_file.load_key(password=password)
        decrypted = io.BytesIO()
        try:
            office_file.decrypt(decrypted)
        except msoffcrypto.exceptions.InvalidKeyError:
            continue
        decrypted.seek(0)
        return decrypted
    return None


def _table_rows(presentation) -> Iterator[list[str]]:
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    yield [cell.text.replace("\n", " ").strip() for cell in row.cells]


class KickOffMeeting:
    def __init__(self, file_path: str, passwords: tuple[str, ...] = DEFAULT_PASSWORDS):
        self.file_path = file_path
        self.passwords = passwords
        self.last_error = ""
        self.info: KickOffInfo = {}

    def load_data(self) -> bool:
        name = os.path.basename(self.file_path)
        try:
            deck = _decrypt(self.file_path, self.passwords)
            if deck is None:
                self.last_error = f"Could not open Kick Off Meeting '{name}': none of the known passwords works."
                return False
            self.info = parse_kick_off(_table_rows(Presentation(deck)))
        except Exception as error:
            self.last_error = f"Could not read Kick Off Meeting '{name}': {error}"
            _logger.error(self.last_error)
            return False

        if not self.info:
            self.last_error = f"No engine or ECU information found in Kick Off Meeting '{name}'."
            return False
        return True
